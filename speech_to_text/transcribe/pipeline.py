"""Local Danish ASR + speaker diarization.

ffmpeg -> 16 kHz mono WAV -> faster-whisper (word timestamps) -> pyannote diarization
-> word-to-speaker alignment -> JSON (same schema as the old insanely-fast-whisper
notebook) -> *_edit.docx.

Output schema, written next to each audio file as <name>-transcription.json:
    {"speakers": [{"timestamp": [start, end], "speaker": "SPEAKER_00", "text": "..."}, ...]}
"""
from __future__ import annotations

import inspect
import json
import os
import re
import subprocess
import tempfile
import time
import traceback
import warnings
import wave
from bisect import bisect_right
from pathlib import Path

from .env import GATED_MODELS, Env, bootstrap, fmt_duration, log

DEFAULT_WHISPER_MODEL = "large-v3"
DEFAULT_LANGUAGE = "da"
DEFAULT_DIARIZATION_MODEL = "pyannote/speaker-diarization-3.1"
DEFAULT_EXTENSIONS = ("mp3", "wma", "m4a", "wav", "flac", "ogg", "mp4", "aac", "opus")
JSON_SUFFIX = "-transcription.json"


# --------------------------------------------------------------------------- audio I/O
def to_wav_16k_mono(src: str, dst: str, ffmpeg: str = "ffmpeg") -> None:
    """Any audio -> 16 kHz mono 16-bit PCM WAV. pcm_s16le keeps the file decodable by the
    stdlib ``wave`` module, so reading never depends on torchaudio/torchcodec backends."""
    r = subprocess.run(
        [ffmpeg, "-y", "-loglevel", "error", "-nostdin", "-i", src,
         "-ac", "1", "-ar", "16000", "-vn", "-c:a", "pcm_s16le", "-f", "wav", dst],
        capture_output=True, text=True,
    )
    if r.returncode != 0:
        raise RuntimeError(f"ffmpeg failed on {src}: {r.stderr.strip()[-500:]}")


def audio_duration_seconds(path: str, ffmpeg: str = "ffmpeg") -> float | None:
    """Length of any audio file via ffmpeg's header parse (no decoding, ~50 ms).
    Works with the pip static ffmpeg too, which ships no ffprobe."""
    r = subprocess.run([ffmpeg, "-nostdin", "-hide_banner", "-i", path], capture_output=True, text=True)
    m = re.search(r"Duration:\s*(\d+):(\d+):(\d+(?:\.\d+)?)", r.stderr)
    if not m:
        return None
    h, mi, s = m.groups()
    return int(h) * 3600 + int(mi) * 60 + float(s)


def wav_duration_seconds(wav_path: str) -> float:
    with wave.open(wav_path, "rb") as wf:
        return wf.getnframes() / float(wf.getframerate())


def load_wav_in_memory(wav_path: str) -> dict:
    """16 kHz mono PCM WAV -> {'waveform': (1, time) float32 tensor, 'sample_rate': int},
    which pyannote.audio 3.x and 4.x accept directly (no torchcodec involved)."""
    import numpy as np
    import torch

    with wave.open(wav_path, "rb") as wf:
        sr, n_channels, sampwidth = wf.getframerate(), wf.getnchannels(), wf.getsampwidth()
        raw = wf.readframes(wf.getnframes())
    if sampwidth != 2:
        raise ValueError(f"expected 16-bit PCM, got {sampwidth * 8}-bit")
    data = np.frombuffer(raw, dtype="<i2").astype(np.float32) / 32768.0
    if n_channels > 1:
        data = data.reshape(-1, n_channels).mean(axis=1)
    return {"waveform": torch.from_numpy(np.ascontiguousarray(data)).unsqueeze(0), "sample_rate": sr}


def write_wav(path: str, samples, sample_rate: int, channels: int = 1) -> None:
    """Write float32 samples in [-1, 1] (shape (time,) or (time, channels)) as 16-bit PCM."""
    import numpy as np

    pcm = (np.clip(np.asarray(samples, dtype=np.float32), -1, 1) * 32767).astype("<i2")
    with wave.open(path, "wb") as wf:
        wf.setnchannels(channels)
        wf.setsampwidth(2)
        wf.setframerate(sample_rate)
        wf.writeframes(np.ascontiguousarray(pcm).tobytes())


# --------------------------------------------------------------------------- files
def transcription_json_for(audio_path: str | os.PathLike) -> str:
    return os.path.splitext(str(audio_path))[0] + JSON_SUFFIX


def list_audio_files(folder: str | os.PathLike, extensions=DEFAULT_EXTENSIONS) -> list[str]:
    """Audio files directly in ``folder`` (not recursive), extensions case-insensitive."""
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f"Audio folder does not exist or is not a folder: {folder}")
    wanted = {e.lower().lstrip(".") for e in extensions}
    return sorted(str(p) for p in folder.iterdir()
                  if p.is_file() and p.suffix.lower().lstrip(".") in wanted)


def write_json_atomic(path: str, payload) -> None:
    """Temp file + rename, so an interrupted run never leaves a truncated JSON behind."""
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


# --------------------------------------------------------------------------- alignment
def assign_speakers(words: list[dict], turns: list[tuple]) -> list[dict]:
    """Tag each word with the speaker whose turn covers its midpoint (else the nearest
    turn), then merge consecutive same-speaker words into utterances."""
    if not turns:
        return [{"timestamp": [w["start"], w["end"]], "speaker": "SPEAKER_00",
                 "text": w["text"].strip()} for w in words]

    starts = [t[0] for t in turns]
    tagged = []
    for w in words:
        mid = 0.5 * (w["start"] + w["end"])
        idx = bisect_right(starts, mid) - 1
        candidates = [turns[i] for i in (idx, idx + 1) if 0 <= i < len(turns)] or [turns[0]]
        containing = [t for t in candidates if t[0] <= mid <= t[1]]
        if containing:
            spk = containing[0][2]
        else:
            spk = min(candidates, key=lambda t: min(abs(mid - t[0]), abs(mid - t[1])))[2]
        tagged.append((w, spk))

    utterances: list[dict] = []
    for w, spk in tagged:
        if utterances and utterances[-1]["speaker"] == spk:
            utterances[-1]["timestamp"][1] = w["end"]
            utterances[-1]["text"] += w["text"] if w["text"].startswith(" ") else " " + w["text"]
        else:
            utterances.append({"timestamp": [w["start"], w["end"]], "speaker": spk, "text": w["text"]})
    for u in utterances:
        u["text"] = u["text"].strip()
    return utterances


# --------------------------------------------------------------------------- models
class _StepLogHook:
    """pyannote progress hook that prints one plain line per pipeline step.

    pyannote's own ProgressHook uses a rich "live" display, which in Jupyter redraws the
    whole cell output and wipes everything printed before it. Plain lines never do that.
    """

    def __init__(self):
        self.step = None
        self.t = None
        self.total = None

    def __call__(self, step_name, step_artifact, file=None, total=None, completed=None):
        if step_name != self.step:
            self.finish()
            self.step, self.t, self.total = step_name, time.time(), total
            log(f"      {step_name} ...")

    def finish(self):
        if self.step is not None:
            log(f"      {self.step} done in {fmt_duration(time.time() - self.t)}")
            self.step = None


class Transcriber:
    """Loads the ASR + diarization models once and transcribes files/folders.

    Both loaders verify themselves: the ASR model runs a real 1 s decode after loading
    and drops to CPU int8 if the GPU path is broken; the diarizer turns the usual pyannote
    failures (terms not accepted / bad token) into an explicit message.
    """

    def __init__(self, env: Env | None = None, whisper_model: str = DEFAULT_WHISPER_MODEL,
                 language: str = DEFAULT_LANGUAGE,
                 diarization_model: str = DEFAULT_DIARIZATION_MODEL,
                 num_speakers: int | None = None):
        self.env = env or bootstrap()
        self.whisper_model = whisper_model
        self.language = language
        self.diarization_model = diarization_model
        self.num_speakers = num_speakers
        warnings.filterwarnings("ignore", message=".*torchcodec.*")  # we never use it
        warnings.filterwarnings("ignore", message=".*TensorFloat-32.*")
        warnings.filterwarnings("ignore", message=".*degrees of freedom.*")  # pyannote on very short audio
        self.asr, self.asr_device, self.asr_compute_type = self._load_asr()
        self.diarizer, self.diar_device = self._load_diarizer()
        log(f"Models ready: {whisper_model} on {self.asr_device}/{self.asr_compute_type}, "
            f"{diarization_model} on {self.diar_device}.")

    def _load_asr(self):
        import numpy as np
        from faster_whisper import WhisperModel

        attempts = [(self.env.asr_device, self.env.asr_compute_type)]
        if self.env.asr_device == "cuda":
            attempts.append(("cpu", "int8"))
        last = None
        for dev, ct in attempts:
            log(f"Loading faster-whisper {self.whisper_model} on {dev} ({ct}) ...")
            t0 = time.time()
            try:
                model = WhisperModel(self.whisper_model, device=dev, compute_type=ct,
                                     download_root=self.env.model_cache_str)
                noise = (np.random.default_rng(0).standard_normal(16000) * 0.01).astype(np.float32)
                segments, _ = model.transcribe(noise, language=self.language, vad_filter=False,
                                               beam_size=1, without_timestamps=True)
                list(segments)  # cuBLAS is loaded lazily: force a real decode
                log(f"   ASR model ready in {fmt_duration(time.time() - t0)}.")
                return model, dev, ct
            except Exception as e:
                last = e
                log(f"  !! failed on {dev}/{ct}: {type(e).__name__}: {str(e)[:300]}")
        raise RuntimeError(f"faster-whisper could not run {self.whisper_model} on any device") from last

    def _load_diarizer(self):
        import torch
        from pyannote.audio import Pipeline

        params = inspect.signature(Pipeline.from_pretrained).parameters
        kwargs = {"token": self.env.hf_token} if "token" in params else {"use_auth_token": self.env.hf_token}
        if "cache_dir" in params and self.env.model_cache_str:
            kwargs["cache_dir"] = self.env.model_cache_str
        hint = ("Check that the HF token is valid and that you accepted the terms of "
                + " and ".join(f"https://hf.co/{m}" for m in GATED_MODELS)
                + " with the account that owns the token.")
        log(f"Loading diarization pipeline {self.diarization_model} ...")
        t0 = time.time()
        try:
            pipeline = Pipeline.from_pretrained(self.diarization_model, **kwargs)
        except Exception as e:
            raise RuntimeError(f"Could not download {self.diarization_model}: "
                               f"{type(e).__name__}: {str(e)[:300]}\n{hint}") from e
        if pipeline is None:
            raise RuntimeError(f"pyannote returned None for {self.diarization_model}. {hint}")
        device = self.env.diar_device
        try:
            pipeline.to(torch.device(device))
        except Exception as e:
            log(f"  !! could not move diarizer to {device} ({e}); using CPU")
            pipeline.to(torch.device("cpu"))
            device = "cpu"
        log(f"   diarization pipeline ready on {device} in {fmt_duration(time.time() - t0)}.")
        return pipeline, device

    # ----------------------------------------------------------------- steps
    def run_whisper(self, wav_path: str, vad_filter: bool = True, progress: bool = True) -> list[dict]:
        """List of {'start','end','text'} words. Shows a progress bar over the audio
        timeline while Whisper works (segments are produced lazily, in order)."""
        segments, _ = self.asr.transcribe(wav_path, language=self.language, task="transcribe",
                                          vad_filter=vad_filter, word_timestamps=True, beam_size=5)
        total = wav_duration_seconds(wav_path)
        bar = None
        if progress and total > 30:
            from tqdm.auto import tqdm
            bar = tqdm(total=round(total), unit="s", desc="   transcribing (audio seconds)", leave=False,
                       bar_format="{desc}: {percentage:3.0f}%|{bar}| {n_fmt}/{total_fmt}s [{elapsed}<{remaining}]")
        words = []
        done = 0.0
        try:
            for seg in segments:
                if bar is not None:
                    bar.update(max(0.0, min(float(seg.end), total) - done))
                    done = max(done, min(float(seg.end), total))
                if not seg.words:
                    words.append({"start": float(seg.start), "end": float(seg.end), "text": seg.text.strip()})
                    continue
                for w in seg.words:
                    if w.start is not None and w.end is not None:
                        words.append({"start": float(w.start), "end": float(w.end), "text": w.word})
        finally:
            if bar is not None:
                bar.update(max(0.0, total - done))
                bar.close()
        return words

    def _diarize(self, audio: dict, num_speakers: int | None, progress: bool):
        kwargs = {"num_speakers": num_speakers} if num_speakers else {}
        if progress and audio["waveform"].shape[1] > 30 * audio["sample_rate"]:
            hook = _StepLogHook()
            try:
                return self.diarizer(audio, hook=hook, **kwargs)
            finally:
                hook.finish()
        return self.diarizer(audio, **kwargs)

    def run_diarization(self, wav_path: str, num_speakers: int | None | str = "default",
                        progress: bool = True) -> list[tuple]:
        """Sorted (start, end, speaker) turns. If forcing the speaker count fails (a file
        with fewer voices than requested), retry with auto-detection."""
        if num_speakers == "default":
            num_speakers = self.num_speakers
        audio = load_wav_in_memory(wav_path)  # in-memory: pyannote never touches torchcodec
        try:
            result = self._diarize(audio, num_speakers, progress)
        except Exception as e:
            if not num_speakers:
                raise
            log(f"  diarization with num_speakers={num_speakers} failed "
                f"({type(e).__name__}: {str(e)[:120]}); retrying with auto-detection")
            result = self._diarize(load_wav_in_memory(wav_path), None, progress)
        annotation = getattr(result, "speaker_diarization", result)  # pyannote 4.x wrapper / 3.x Annotation
        turns = [(float(t.start), float(t.end), str(spk)) for t, _, spk in annotation.itertracks(yield_label=True)]
        turns.sort()
        return turns

    def transcribe_file(self, audio_path: str | os.PathLike, out_json: str | None = None,
                        vad_filter: bool = True, progress: bool = True) -> str:
        """Full pipeline for one file. Writes <name>-transcription.json next to the audio
        (or to ``out_json``) and returns that path. Logs every step with its duration.
        ``self.last_stats`` holds timing info for the last file."""
        audio_path = str(audio_path)
        out_json = out_json or transcription_json_for(audio_path)
        t0 = time.time()
        with tempfile.TemporaryDirectory() as tmp:
            wav = os.path.join(tmp, "audio.wav")
            to_wav_16k_mono(audio_path, wav, self.env.ffmpeg)
            audio_len = wav_duration_seconds(wav)
            t1 = time.time()
            if progress:
                log(f"   audio length {fmt_duration(audio_len)}; converted in {fmt_duration(t1 - t0)}. "
                    f"Transcribing with {self.whisper_model} on {self.asr_device} ...")
            words = self.run_whisper(wav, vad_filter=vad_filter, progress=progress)
            t2 = time.time()
            if progress:
                log(f"   transcribed {len(words)} words in {fmt_duration(t2 - t1)} "
                    f"({audio_len / max(t2 - t1, 1e-6):.1f}x realtime). Finding speakers ...")
            turns = self.run_diarization(wav, progress=progress)
            t3 = time.time()
        utterances = assign_speakers(words, turns)
        write_json_atomic(out_json, {"speakers": utterances})
        speakers = sorted({u["speaker"] for u in utterances})
        total = time.time() - t0
        self.last_stats = {"audio_seconds": audio_len, "seconds": total, "words": len(words),
                           "speakers": len(speakers)}
        if progress:
            log(f"   speakers found: {len(speakers)} in {fmt_duration(t3 - t2)}. File done in "
                f"{fmt_duration(total)} ({audio_len / max(total, 1e-6):.1f}x realtime) -> {os.path.basename(out_json)}")
        return out_json

    def transcribe_folder(self, folder: str | os.PathLike, extensions=DEFAULT_EXTENSIONS,
                          skip_existing: bool = True) -> dict:
        """Transcribe every audio file in ``folder``. A failing file is reported and does
        not stop the batch. Prints the plan up front and an estimated remaining time after
        each file. Returns {'done': [...], 'skipped': [...], 'failed': [(file, err)]}."""
        import torch

        files = list_audio_files(folder, extensions)
        todo = [f for f in files if not (skip_existing and os.path.exists(transcription_json_for(f)))]
        skipped = [f for f in files if f not in todo]
        durations = {f: (audio_duration_seconds(f, self.env.ffmpeg) or 0.0) for f in todo}
        total_audio = sum(durations.values())
        log(f"{len(files)} audio file(s) in {folder}: {len(todo)} to transcribe "
            f"({fmt_duration(total_audio)} of audio), {len(skipped)} already done (skipped).")
        for f in todo:
            log(f"   todo  {os.path.basename(f)}  ({fmt_duration(durations[f])})")
        if todo:
            log("The first file tells us the speed; an estimate of the remaining time follows after it.")

        done, failed = [], []
        processed_audio = processed_seconds = 0.0
        batch_t0 = time.time()
        for idx, audio in enumerate(todo, start=1):
            remaining_audio = sum(durations[f] for f in todo[idx - 1:])
            if processed_audio:
                speed = processed_audio / max(processed_seconds, 1e-6)
                estimate = f"about {fmt_duration(remaining_audio / speed)} left for this and the remaining files"
            else:
                estimate = "estimate follows after this file"
            log(f"=== FILE {idx} of {len(todo)}: {os.path.basename(audio)} ({fmt_duration(durations[audio])}) "
                f"| elapsed {fmt_duration(time.time() - batch_t0)} | {estimate} ===")
            try:
                self.transcribe_file(audio)
                done.append(audio)
                processed_audio += self.last_stats["audio_seconds"]
                processed_seconds += self.last_stats["seconds"]
            except Exception as e:
                log(f"   !! failed: {type(e).__name__}: {e}")
                traceback.print_exc(limit=3)
                failed.append((audio, f"{type(e).__name__}: {e}"))
            finally:
                if torch.cuda.is_available():
                    torch.cuda.empty_cache()
            remaining_audio = sum(durations[f] for f in todo[idx:])
            if remaining_audio and processed_audio:
                speed = processed_audio / max(processed_seconds, 1e-6)
                log(f"--- PROGRESS: {idx} of {len(todo)} files done, {len(todo) - idx} to go "
                    f"({fmt_duration(remaining_audio)} of audio); at {speed:.1f}x realtime "
                    f"that is about {fmt_duration(remaining_audio / speed)} more. ---")
        log(f"Batch finished in {fmt_duration(time.time() - batch_t0)}: {len(done)} transcribed, "
            f"{len(skipped)} skipped, {len(failed)} failed.")
        return {"done": done, "skipped": skipped, "failed": failed}


# --------------------------------------------------------------------------- docx
def diarization_to_docx_edit(json_path: str) -> str:
    """Same *_edit.docx layout as the old notebook: centered bold file name, then one
    paragraph per utterance with timestamp, speaker and text on separate lines."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run(os.path.basename(json_path)).bold = True
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    for entry in data["speakers"]:
        paragraph = doc.add_paragraph()
        paragraph.add_run(str(entry["timestamp"]) + "\n")
        paragraph.add_run(str(entry["speaker"]) + "\n")
        paragraph.add_run(str(entry["text"]).strip())
    docx_path = json_path[: -len(".json")] + "_edit.docx"
    doc.save(docx_path)
    return docx_path


def convert_folder_to_docx(folder: str | os.PathLike) -> list[str]:
    files_json = sorted(str(p) for p in Path(folder).glob(f"*{JSON_SUFFIX}"))
    out = []
    log(f"{len(files_json)} transcription JSON file(s) -> docx")
    for path in files_json:
        try:
            out.append(diarization_to_docx_edit(path))
        except Exception as e:
            log(f"   !! {os.path.basename(path)}: {type(e).__name__}: {e}")
    return out


# --------------------------------------------------------------------------- one call
def run(path_audio: str | os.PathLike, num_speakers: int | None = 2, skip_existing: bool = True,
        extensions=DEFAULT_EXTENSIONS, whisper_model: str = DEFAULT_WHISPER_MODEL,
        language: str = DEFAULT_LANGUAGE, selftest: bool = True, docx: bool = True,
        hf_token: str | None = None) -> dict:
    """Everything in one call: prepare the environment, load models, self-test, transcribe
    the folder, convert to docx. This is what the notebook calls. ``hf_token`` is optional;
    when given it is used and saved to the repo .env."""
    t0 = time.time()
    log("STEP 1/4  Preparing the environment (packages, GPU, ffmpeg, token) ...")
    env = bootstrap(hf_token=hf_token)
    log("STEP 2/4  Loading models (the first time on a machine this downloads ~3 GB, 1-5 min; "
        "afterwards ~10 s) ...")
    transcriber = Transcriber(env, whisper_model=whisper_model, language=language,
                              num_speakers=num_speakers)
    if selftest:
        log("STEP 3/4  Self-test on a 3 s tone (~5 s) ...")
        from .selftest import smoke_test
        smoke_test(transcriber)
    log(f"STEP 4/4  Transcribing {path_audio} ...")
    summary = transcriber.transcribe_folder(path_audio, extensions, skip_existing=skip_existing)
    if docx:
        summary["docx"] = convert_folder_to_docx(path_audio)
    if summary["failed"]:
        log("Failed files:")
        for a, e in summary["failed"]:
            log(f"   {a} -> {e}")
    else:
        log(f"ALL DONE in {fmt_duration(time.time() - t0)}. Results are next to the audio files.")
    return summary
